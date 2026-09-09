"""체크시트 xlsx 생성 테스트 — 템플릿 zip 편집, 이미지 보존, 셀 값."""
from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
import pytest

from src.core.inspection.check_sheet import (
    check_sheet_name,
    render_sheet_xml,
    write_check_sheet,
)
from src.core.inspection.lot_builder import CheckSheetSpec, build_lots
from src.core.inspection.mtc_parser import load_mtc_run

FIXTURE = Path(__file__).parent / "fixtures" / "mtc"
TEMPLATE = FIXTURE / "check_sheet_template.xlsx"


def test_name():
    assert check_sheet_name("P2401002", 12, "AC160TS") == "P2401002_12M_AC160TS.xlsx"


def test_write_check_sheet_fills_cells_and_keeps_images(tmp_path):
    out = tmp_path / "P2401002_12M_AC160TS.xlsx"
    freqs = [284.81, 259.42, 364.66] + [None] * 9
    qs = [422.81, 141.68, 79.18] + [None] * 9
    checks = {"backside": None, "a_plus_b": True, "unipeak": False, "noise": True, "frequency": False}
    write_check_sheet(TEMPLATE, out, "P2401002", "AC160TS", freqs, qs, checks)

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert "xl/media/image1.png" in names and "xl/drawings/drawing1.xml" in names   # 이미지 보존
        assert z.read("xl/media/image1.png") == zipfile.ZipFile(TEMPLATE).read("xl/media/image1.png")

    ws = openpyxl.load_workbook(out, data_only=True).worksheets[0]
    assert ws["B4"].value == "P2401002" and ws["B5"].value == "AC160TS"
    assert ws["L26"].value == "■" and ws["M26"].value == "□"      # backside: 템플릿 유지
    assert ws["L28"].value == "■" and ws["M28"].value == "□"      # A+B pass
    assert ws["L31"].value == "□" and ws["M31"].value == "■"      # unipeak fail
    assert ws["L32"].value == "■" and ws["M32"].value == "□"      # noise pass
    assert ws["L34"].value == "□" and ws["M34"].value == "■"      # frequency fail
    assert ws["B41"].value == 284.81 and ws["C41"].value == 259.42 and ws["D41"].value == 364.66
    assert ws["B42"].value == 422.81 and ws["D42"].value == 79.18
    assert ws["E41"].value is None and ws["M41"].value is None and ws["M42"].value is None
    assert ws["A53"].value.startswith("   Warranty")      # 나머지 텍스트 유지


def test_render_escapes_and_missing_cells():
    xml = '<row><c r="B4" s="11" t="s"><v>23</v></c><c r="B41" s="9" t="n"><v>1</v></c></row>'
    out = render_sheet_xml(xml, "A&B<1>", "T", [12.5], [None], {})
    assert '<c r="B4" s="11" t="inlineStr"><is><t>A&amp;B&lt;1&gt;</t></is></c>' in out
    assert '<c r="B41" s="9"><v>12.5</v></c>' in out
    assert "B5" not in out          # 템플릿에 없는 셀은 건드리지 않음


def test_missing_template_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        write_check_sheet(tmp_path / "nope.xlsx", tmp_path / "out.xlsx", "U", "T", [], [], {})


def test_check_sheet_spec_lot_checks():
    spec = CheckSheetSpec("t.xlsx", "AC160TS", {
        "1101": {"a_plus_b": True, "unipeak": True, "noise": True, "frequency": True},
        "1102": {"a_plus_b": False, "unipeak": True, "noise": False, "frequency": True},
    })
    lc = spec.lot_checks(["1101", "1102"])
    assert lc == {"backside": None, "a_plus_b": False, "unipeak": True, "noise": False, "frequency": True}
    assert spec.lot_checks(["9999"]) == {"backside": None, "a_plus_b": None, "unipeak": None,
                                         "noise": None, "frequency": None}


def test_build_lots_with_check_sheet(tmp_path):
    run = load_mtc_run(FIXTURE / "20260909")
    slots = [s for s in run.slots if s.has_sweep]
    spec = CheckSheetSpec(str(TEMPLATE), "AC160TS",
                          {s.code: {"a_plus_b": True, "unipeak": True, "noise": True, "frequency": True}
                           for s in slots})
    lots, _ = build_lots(tmp_path, slots, [3], "P2401002", "AC160", "b", spec)
    sheet = Path(lots[0].check_sheet)
    assert sheet.name == "P2401002_3M_AC160TS.xlsx" and sheet.parent.name == "P2401002_3M_AC160"
    ws = openpyxl.load_workbook(sheet, data_only=True).worksheets[0]
    assert ws["B4"].value == "P2401002" and ws["B41"].value == 284.81 and ws["E41"].value is None
    # 체크시트 없이도 동작
    lots2, _ = build_lots(tmp_path / "b", slots, [3], "P1", "AC160", "b", None)
    assert lots2[0].check_sheet is None
